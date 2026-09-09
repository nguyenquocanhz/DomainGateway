"""Xuat danh sach ten mien ra Markdown, XLSX va PDF.

Ba dinh dang phuc vu ba muc dich khac nhau:
  - Markdown : cho AI agent doc. Kem ngu canh (nguon du lieu, nguong canh bao,
               bay nghiep vu) chu khong chi la bang so lieu tran.
  - XLSX     : cho nguoi doc va loc. Ngay thang la kieu date that de Excel sap
               xep dung, co autofilter va to mau theo muc do khan.
  - PDF      : de in / gui di. Nhung Vietnamese can font TTF nhung vao file,
               font PDF loi (Helvetica) khong co dau tieng Viet.
"""

from __future__ import annotations

import io
import os
from datetime import datetime, timezone

# Excel va LibreOffice coi o bat dau bang cac ky tu nay la CONG THUC va chay no
# ngay khi mo file. Du lieu o day khong phai do nguoi dung go het: ten registrar
# lay thang tu phan hoi WHOIS/RDAP cua ben thu ba. Da kiem bang openpyxl: chuoi
# "=1+1" ghi ra o co data_type 'f' - cong thuc that, khong phai chu.
KY_TU_CONG_THUC = ("=", "+", "-", "@", "\t", "\r")


def rao_cong_thuc(value):
    """Chen dau nhay don de bang tinh doc thanh chu, khong phai cong thuc.

    Chi dong vao chuoi. So va ngay thang phai giu nguyen kieu, bien thanh chuoi
    la Excel het sap xep dung theo thoi gian - dung cai ma XLSX sinh ra de co.
    """
    if not isinstance(value, str):
        return value
    return "'" + value if value[:1] in KY_TU_CONG_THUC else value


STATUS_LABEL_VI = {
    "active": "Đang hoạt động",
    "expiring": "Sắp hết hạn",
    "critical": "Nguy cấp",
    "expired": "Đã hết hạn",
    "unknown": "Chưa rõ",
}

# Mau theo muc do khan, dung chung cho ca XLSX va PDF
STATUS_COLOR = {
    "active": "1E8449",
    "expiring": "B7791F",
    "critical": "C0392B",
    "expired": "922B21",
    "unknown": "7F8C8D",
}

# Font co dau tieng Viet. Tim theo thu tu, dung cai dau tien co that.
FONT_CANDIDATES = [
    (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
    (r"C:\Windows\Fonts\segoeui.ttf", r"C:\Windows\Fonts\segoeuib.ttf"),
    (r"C:\Windows\Fonts\tahoma.ttf", r"C:\Windows\Fonts\tahomabd.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
]


def _local(dt):
    """Doi sang gio may chu (= gio nguoi dung) roi bo tzinfo cho Excel."""
    return dt.astimezone().replace(tzinfo=None) if dt else None


def _fmt(dt) -> str:
    return dt.astimezone().strftime("%d/%m/%Y") if dt else "—"


def _days(rec) -> str:
    d = rec.days_left
    if d is None:
        return "—"
    return f"quá {abs(d)} ngày" if d < 0 else f"{d} ngày"


def _sorted(records):
    return sorted(records, key=lambda r: (r.days_left is None,
                                          r.days_left if r.days_left is not None else 0))


def _summary(records, warn: int, crit: int) -> dict:
    buckets = {k: 0 for k in STATUS_LABEL_VI}
    for rec in records:
        buckets[rec.status(warn, crit)] += 1
    providers = sorted({(r.provider or r.registrar or "Chưa rõ") for r in records})
    return {
        "total": len(records),
        "buckets": buckets,
        "providers": providers,
        "need_action": buckets["expiring"] + buckets["critical"] + buckets["expired"],
    }


# --------------------------------------------------------------------------- #
# Markdown — soan cho AI agent doc
# --------------------------------------------------------------------------- #
def _md_cell(value) -> str:
    """Bang Markdown dung | lam dau cot nen phai escape, va bo xuong dong."""
    text = str(value if value is not None else "").replace("|", r"\|")
    return " ".join(text.split()) or "—"


def to_markdown(records, warn: int = 45, crit: int = 14, title: str = None) -> str:
    records = _sorted(records)
    info = _summary(records, warn, crit)
    now = datetime.now(timezone.utc).astimezone()

    out = [f"# {title or 'Danh mục tên miền'}", ""]
    out += [
        "> **Ngữ cảnh cho agent** — đọc phần này trước khi kết luận.",
        f"> - Xuất lúc: `{now.strftime('%Y-%m-%d %H:%M:%S %z')}` (giờ địa phương)",
        "> - Nguồn: đọc trực tiếp từ registry qua RDAP, WHOIS cổng 43 và API VNNIC/BKNS.",
        f"> - Ngưỡng: *sắp hết hạn* ≤ **{warn}** ngày, *nguy cấp* ≤ **{crit}** ngày.",
        "> - `Còn lại` tính từ thời điểm xuất file ở trên, không phải thời điểm bạn đọc.",
        "",
        "## Tóm tắt",
        "",
        f"- Tổng cộng **{info['total']}** tên miền ở **{len(info['providers'])}** nhà cung cấp.",
        f"- Cần xử lý: **{info['need_action']}** "
        f"(sắp hết hạn {info['buckets']['expiring']}, "
        f"nguy cấp {info['buckets']['critical']}, "
        f"đã hết hạn {info['buckets']['expired']}).",
        f"- Không đọc được ngày hết hạn: **{info['buckets']['unknown']}**.",
        "",
    ]

    due = [r for r in records if r.status(warn, crit) != "active"
           and r.status(warn, crit) != "unknown"]
    if due:
        out += ["## Cần hành động", "",
                "| Tên miền | Trạng thái | Còn lại | Ngày hết hạn | Nhà cung cấp |",
                "|---|---|---:|---|---|"]
        for rec in due:
            out.append("| `{}` | {} | {} | {} | {} |".format(
                _md_cell(rec.domain), _md_cell(STATUS_LABEL_VI[rec.status(warn, crit)]),
                _md_cell(_days(rec)), _md_cell(_fmt(rec.expires_at)),
                _md_cell(rec.provider or rec.registrar)))
        out.append("")

    out += ["## Toàn bộ tên miền", "",
            "| Tên miền | Trạng thái | Còn lại | Ngày hết hạn | Ngày đăng ký | "
            "Nhà cung cấp | Registrar | Nameserver | Tag | Nguồn |",
            "|---|---|---:|---|---|---|---|---|---|---|"]
    for rec in records:
        out.append("| `{}` | {} | {} | {} | {} | {} | {} | {} | {} | {} |".format(
            _md_cell(rec.domain),
            _md_cell(STATUS_LABEL_VI[rec.status(warn, crit)]),
            _md_cell(_days(rec)),
            _md_cell(_fmt(rec.expires_at)),
            _md_cell(_fmt(rec.created_at)),
            _md_cell(rec.provider),
            _md_cell(rec.registrar),
            _md_cell(", ".join(rec.nameservers)),
            _md_cell(", ".join(rec.tags)),
            _md_cell(rec.source),
        ))
    out.append("")

    failed = [r for r in records if r.error]
    if failed:
        out += ["## Tên miền tra cứu lỗi", ""]
        for rec in failed:
            out.append(f"- `{_md_cell(rec.domain)}` — {_md_cell(rec.error)}")
        out.append("")

    out += [
        "## Lưu ý nghiệp vụ",
        "",
        "- **Ngày hết hạn không phải ngày mất tên miền.** Sau khi hết hạn còn "
        "Auto-Renew Grace Period (0–45 ngày, gia hạn giá thường), rồi Redemption "
        "(30 ngày, phí chuộc thường 80–200 USD), rồi Pending Delete (5 ngày). "
        "Nhưng **website tắt ngay khi hết hạn**, không đợi hết grace period.",
        "- **`clientTransferProhibited` là bình thường**, đó là khoá chống chuyển "
        "nhượng trái phép mà đa số registrar bật mặc định — không phải sự cố.",
        "- **Cột `Nguồn`**: `rdap` là chuẩn ICANN, `bkns` là API WHOIS Việt Nam, "
        "`whois43` là WHOIS thô, `manual` là ngày do người dùng tự nhập.",
        "- Registry `.eu` không công bố ngày hết hạn nên các tên miền `.eu` "
        "thường phải nhập tay.",
        "",
    ]
    return "\n".join(out)


# --------------------------------------------------------------------------- #
# XLSX
# --------------------------------------------------------------------------- #
def to_xlsx(records, warn: int = 45, crit: int = 14) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    records = _sorted(records)
    info = _summary(records, warn, crit)

    wb = Workbook()
    ws = wb.active
    ws.title = "Tên miền"

    headers = ["Tên miền", "Trạng thái", "Còn lại (ngày)", "Ngày hết hạn", "Ngày đăng ký",
               "Nhà cung cấp", "Registrar", "Nameserver", "Tự động gia hạn", "Tag",
               "Nguồn", "Ghi chú", "Lỗi"]
    ws.append(headers)

    head_fill = PatternFill("solid", fgColor="1F2933")
    head_font = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="D5D8DC")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)

    for rec in records:
        status = rec.status(warn, crit)
        ws.append([rao_cong_thuc(o) for o in (
            rec.domain,
            STATUS_LABEL_VI[status],
            rec.days_left,                      # so that -> Excel sap xep/loc duoc
            _local(rec.expires_at),             # date that, khong phai chuoi
            _local(rec.created_at),
            rec.provider,
            rec.registrar,
            ", ".join(rec.nameservers),
            "" if rec.auto_renew is None else ("Có" if rec.auto_renew else "Không"),
            ", ".join(rec.tags),
            rec.source,
            rec.note,
            rec.error,
        )])
        row = ws.max_row
        ws.cell(row=row, column=2).font = Font(bold=status != "active",
                                               color=STATUS_COLOR[status])
        ws.cell(row=row, column=3).alignment = Alignment(horizontal="right")
        for col in (4, 5):
            ws.cell(row=row, column=col).number_format = "DD/MM/YYYY"

    widths = [26, 16, 14, 14, 14, 22, 30, 40, 15, 14, 10, 28, 34]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    # ---- sheet tom tat ----
    s = wb.create_sheet("Tóm tắt")
    now = datetime.now(timezone.utc).astimezone()
    rows = [
        ("Xuất lúc", now.strftime("%d/%m/%Y %H:%M")),
        ("Tổng số tên miền", info["total"]),
        ("Số nhà cung cấp", len(info["providers"])),
        ("", ""),
        (f"Sắp hết hạn (≤ {warn} ngày)", info["buckets"]["expiring"]),
        (f"Nguy cấp (≤ {crit} ngày)", info["buckets"]["critical"]),
        ("Đã hết hạn", info["buckets"]["expired"]),
        ("Đang hoạt động", info["buckets"]["active"]),
        ("Chưa đọc được ngày hết hạn", info["buckets"]["unknown"]),
        ("", ""),
        ("Nguồn dữ liệu", "RDAP · WHOIS:43 · API VNNIC/BKNS"),
    ]
    for label, value in rows:
        s.append([label, value])
    for row in range(1, s.max_row + 1):
        s.cell(row=row, column=1).font = Font(bold=True)
    s.column_dimensions["A"].width = 32
    s.column_dimensions["B"].width = 36

    s.append([])
    s.append(["Nhà cung cấp", "Số tên miền"])
    s.cell(row=s.max_row, column=1).font = Font(bold=True, color="FFFFFF")
    s.cell(row=s.max_row, column=2).font = Font(bold=True, color="FFFFFF")
    s.cell(row=s.max_row, column=1).fill = head_fill
    s.cell(row=s.max_row, column=2).fill = head_fill
    counts = {}
    for rec in records:
        key = rec.provider or rec.registrar or "Chưa rõ"
        counts[key] = counts.get(key, 0) + 1
    for name, num in sorted(counts.items(), key=lambda kv: -kv[1]):
        s.append([name, num])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _register_font():
    """Nhung font TTF co dau tieng Viet. Tra ve (ten_thuong, ten_dam)."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for regular, bold in FONT_CANDIDATES:
        if os.path.exists(regular) and os.path.exists(bold):
            try:
                pdfmetrics.registerFont(TTFont("DGSans", regular))
                pdfmetrics.registerFont(TTFont("DGSans-Bold", bold))
                return "DGSans", "DGSans-Bold"
            except Exception:
                continue
    # Khong tim thay font nao: Helvetica se lam mat dau tieng Viet, nhung van
    # con hon la khong xuat duoc file.
    return "Helvetica", "Helvetica-Bold"


def to_pdf(records, warn: int = 45, crit: int = 14, title: str = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)

    font, font_bold = _register_font()
    records = _sorted(records)
    info = _summary(records, warn, crit)
    now = datetime.now(timezone.utc).astimezone()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=14 * mm, rightMargin=14 * mm,
        topMargin=14 * mm, bottomMargin=14 * mm,
        title=title or "Danh mục tên miền", author="Domain Gateway",
    )

    h1 = ParagraphStyle("h1", fontName=font_bold, fontSize=16, leading=20,
                        textColor=colors.HexColor("#1F2933"))
    meta = ParagraphStyle("meta", fontName=font, fontSize=8.5, leading=12,
                          textColor=colors.HexColor("#6B7280"))
    cell = ParagraphStyle("cell", fontName=font, fontSize=8, leading=10)
    cell_head = ParagraphStyle("cellhead", fontName=font_bold, fontSize=8.5,
                               leading=11, textColor=colors.white)

    story = [
        Paragraph(title or "Danh mục tên miền", h1),
        Spacer(1, 3 * mm),
        Paragraph(
            f"Xuất lúc {now.strftime('%d/%m/%Y %H:%M')} · "
            f"{info['total']} tên miền · {len(info['providers'])} nhà cung cấp · "
            f"cần xử lý {info['need_action']} · "
            f"ngưỡng cảnh báo {warn} ngày, nguy cấp {crit} ngày<br/>"
            f"Nguồn: RDAP · WHOIS cổng 43 · API VNNIC/BKNS",
            meta),
        Spacer(1, 5 * mm),
    ]

    headers = ["Tên miền", "Trạng thái", "Còn lại", "Ngày hết hạn",
               "Nhà cung cấp", "Registrar", "Nameserver"]
    data = [[Paragraph(h, cell_head) for h in headers]]
    for rec in records:
        status = rec.status(warn, crit)
        colour = STATUS_COLOR[status]
        data.append([
            Paragraph(rec.domain, cell),
            Paragraph(f'<font color="#{colour}">{STATUS_LABEL_VI[status]}</font>', cell),
            Paragraph(_days(rec), cell),
            Paragraph(_fmt(rec.expires_at), cell),
            Paragraph(rec.provider or "—", cell),
            Paragraph(rec.registrar or "—", cell),
            Paragraph("<br/>".join(rec.nameservers) or "—", cell),
        ])

    table = Table(data, repeatRows=1, hAlign="LEFT",
                  colWidths=[46 * mm, 26 * mm, 20 * mm, 24 * mm, 46 * mm, 46 * mm, 52 * mm])
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2933")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#E5E7EB")),
        ("ALIGN", (2, 1), (2, -1), "RIGHT"),
    ]
    for idx in range(1, len(data)):
        if idx % 2 == 0:
            style.append(("BACKGROUND", (0, idx), (-1, idx), colors.HexColor("#FAFAFB")))
    table.setStyle(TableStyle(style))
    story.append(table)

    def _footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#9CA3AF"))
        canvas.drawString(14 * mm, 8 * mm, "Domain Gateway")
        canvas.drawRightString(landscape(A4)[0] - 14 * mm, 8 * mm,
                               f"Trang {canvas.getPageNumber()}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# PDF danh ba nha dang ky
# --------------------------------------------------------------------------- #
def registry_pdf(payload: dict) -> bytes:
    """Dung PDF danh ba tu phan DA LOC do client gui len.

    Bo loc (bo dau tieng Viet, khop danh tinh vs dien giai) nam o JavaScript.
    Neu server tu loc lai thi phai viet logic do lan thu hai bang Python, va hai
    ban se lech nhau. Nen server chi nhan ket qua va lo phan dan trang.

    payload la du lieu tu trinh duyet -> escape truoc khi dua vao Paragraph,
    vi Paragraph cua reportlab hieu mot tap the giong HTML.
    """
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (KeepTogether, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    font, font_bold = _register_font()
    sections = payload.get("sections") or []
    query = str(payload.get("query") or "").strip()
    region = str(payload.get("region_label") or "").strip()
    now = datetime.now(timezone.utc).astimezone()
    total = sum(len(s.get("items") or []) for s in sections)

    def par(text, style):
        return Paragraph(escape(str(text if text is not None else "")), style)

    # par() da tu str() moi thu, nhung vai cho ben duoi noi chuoi TRUOC khi
    # goi par() - va payload den tu trinh duyet nen mot truong sai kieu la
    # TypeError, tra 500 cho mot loi cua nguoi gui. Nen moi phep noi o day
    # deu str() tuong minh. Va tai day chu khong o tang API: them mot truong
    # moi vao PDF thi khong phai nho di sua bo kiem ben app.py.

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=14 * mm, rightMargin=14 * mm,
        topMargin=14 * mm, bottomMargin=14 * mm,
        title="Danh bạ nhà đăng ký tên miền", author="Domain Gateway",
    )

    h1 = ParagraphStyle("h1", fontName=font_bold, fontSize=16, leading=20,
                        textColor=colors.HexColor("#1F2933"))
    h2 = ParagraphStyle("h2", fontName=font_bold, fontSize=12, leading=16,
                        textColor=colors.HexColor("#1F2933"), spaceBefore=2)
    meta = ParagraphStyle("meta", fontName=font, fontSize=8.5, leading=12,
                          textColor=colors.HexColor("#6B7280"))
    note = ParagraphStyle("note", fontName=font, fontSize=8, leading=11,
                          textColor=colors.HexColor("#8A5200"))
    cell = ParagraphStyle("cell", fontName=font, fontSize=8, leading=10.5)
    cell_sm = ParagraphStyle("cellsm", fontName=font, fontSize=7, leading=9.5,
                             textColor=colors.HexColor("#6B7280"))
    cell_head = ParagraphStyle("cellhead", fontName=font_bold, fontSize=8.5,
                               leading=11, textColor=colors.white)

    scope = []
    if query:
        scope.append(f'từ khoá "{query}"')
    if region:
        scope.append(f"khu vực {region}")

    story = [
        par("Danh bạ nhà đăng ký tên miền", h1),
        Spacer(1, 3 * mm),
        par(f"Xuất lúc {now.strftime('%d/%m/%Y %H:%M')} · {total} nhà đăng ký · "
            f"{len(sections)} khu vực"
            + (" · Lọc theo " + ", ".join(scope) if scope else ""), meta),
        Spacer(1, 6 * mm),
    ]

    head_bg = colors.HexColor("#1F2933")
    for sec in sections:
        block = [par(f"{sec.get('code', '')}  {sec.get('name', '')}", h2), Spacer(1, 2 * mm)]

        lines = []
        if sec.get("registry"):
            lines.append(f"Registry: {sec['registry']}")
        if sec.get("whois"):
            lines.append(f"WHOIS:43: {sec['whois']}")
        if sec.get("tlds"):
            lines.append("Đuôi tên miền: " + " ".join(str(x) for x in sec["tlds"]))
        if lines:
            block.append(par(" · ".join(lines), meta))
        if sec.get("yeu_cau"):
            block.append(Spacer(1, 1.5 * mm))
            block.append(par("Điều kiện đăng ký: " + str(sec["yeu_cau"]), note))
        if sec.get("canh_bao"):
            block.append(par("Lưu ý: " + str(sec["canh_bao"]), note))
        block.append(Spacer(1, 3 * mm))
        story.append(KeepTogether(block))

        rows = [[par(h, cell_head) for h in
                 ("Nhà đăng ký", "Quốc gia / Loại hình", "Trustee", "API", "Ghi chú")]]
        for item in sec.get("items") or []:
            name = escape(str(item.get("ten") or ""))
            url = escape(str(item.get("url") or ""))
            rows.append([
                Paragraph(f"<b>{name}</b>" + (f"<br/><font size=7 color='#6B7280'>{url}</font>"
                                              if url else ""), cell),
                par(" · ".join(str(x) for x in (item.get("quoc_gia"), item.get("loai")) if x), cell),
                par("Có" if item.get("trustee") else "", cell),
                par(item.get("api") or "", cell_sm),
                par(item.get("ghi_chu") or "", cell_sm),
            ])

        if len(rows) == 1:
            story.append(par("Không có nhà đăng ký nào khớp trong khu vực này.", meta))
        else:
            # Trustee cần đủ rộng cho chính chữ "Trustee" ở tiêu đề (8.5pt đậm)
            # cộng 6pt padding mỗi bên, nếu không tiêu đề bị ngắt thành "Truste/e".
            table = Table(rows, repeatRows=1, hAlign="LEFT",
                          colWidths=[46 * mm, 28 * mm, 18 * mm, 44 * mm, 46 * mm])
            style = [
                ("BACKGROUND", (0, 0), (-1, 0), head_bg),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#E5E7EB")),
            ]
            for idx in range(1, len(rows)):
                if idx % 2 == 0:
                    style.append(("BACKGROUND", (0, idx), (-1, idx),
                                  colors.HexColor("#FAFAFB")))
            table.setStyle(TableStyle(style))
            story.append(table)
        story.append(Spacer(1, 8 * mm))

    if not sections:
        story.append(par("Không có nhà đăng ký nào khớp bộ lọc hiện tại.", meta))

    def _footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#9CA3AF"))
        canvas.drawString(14 * mm, 8 * mm, "Domain Gateway · danh bạ nhà đăng ký")
        canvas.drawRightString(A4[0] - 14 * mm, 8 * mm, f"Trang {canvas.getPageNumber()}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()
